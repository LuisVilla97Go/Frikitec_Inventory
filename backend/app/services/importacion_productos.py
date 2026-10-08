import io
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any
from uuid import UUID
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.xml import DEFUSEDXML
from pydantic import ValidationError
from app.extensions import db
from app.models.almacen import DimAlmacen
from app.models.producto import CategoriaProducto, DimProducto
from app.models.stock_almacen import FactStockAlmacen
from app.schemas.producto_schema import (
    COLUMNA_DE_CAMPO,
    MAX_STOCK_INICIAL,
    CargaInicial,
    FilaDeImportacion,
    ProductoCrear,
)
from app.services import productos_service
from app.services.errores import DatosInvalidos

MAX_FILAS = 2000
MAX_DESCOMPRIMIDO = 20 * 1024 * 1024
MAX_PARTES_ZIP = 100
MAX_INCIDENCIAS = 300
FIRMA_ZIP = b"PK\x03\x04"

COLUMNAS = {
    "Categoria": "category",
    "Sub Categoría": "sub_category",
    "Producto": "name",
    "Variante": "variant",
    "Código Barras": "barcode",
    "SKU": "sku",
    "Marca": "brand",
    "Costo Unitario": "purchase_price",
    "Precio Venta": "sale_price",
}
OPCIONALES = {"Precio Venta"}
PREFIJO_STOCK = "Stock "
STOCK_TOTAL = "Stock Total"
DERIVADAS = {STOCK_TOTAL, "Costo Total"}

_ENCABEZADO_DE_CAMPO = {campo: encabezado for encabezado, campo in COLUMNAS.items()}
_DINERO = {"purchase_price", "sale_price"}


def _clave(texto: str) -> str:
    """«Sub Categoría», «sub categoria » y «SUB CATEGORIA» son la misma columna."""
    sin_tildes = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return " ".join(sin_tildes.lower().split())


_CATEGORIAS = {_clave(c.value): c.value for c in CategoriaProducto}


@dataclass(frozen=True)
class Incidencia:
    fila: int | None
    columna: str | None
    mensaje: str


@dataclass
class FilaLeida:
    numero: int
    datos: FilaDeImportacion
    stock: dict[str, int]


@dataclass
class Libro:
    almacenes: list[str]
    filas: list[FilaLeida] = field(default_factory=list)
    errores: list[Incidencia] = field(default_factory=list)
    avisos: list[Incidencia] = field(default_factory=list)


@dataclass
class _Alta:
    fila: FilaLeida
    datos: ProductoCrear


@dataclass
class _Cambio:
    fila: FilaLeida
    producto: DimProducto
    campos: dict[str, Any]
    cargar_stock: bool


@dataclass
class Revision:
    errores: list[Incidencia]
    avisos: list[Incidencia]
    almacenes: dict[str, DimAlmacen] = field(default_factory=dict)
    almacenes_nuevos: list[str] = field(default_factory=list)
    altas: list[_Alta] = field(default_factory=list)
    cambios: list[_Cambio] = field(default_factory=list)
    sin_cambios: int = 0

    def resumen(self) -> dict:
        return {
            "nuevos": len(self.altas),
            "actualizados": len(self.cambios),
            "sin_cambios": self.sin_cambios,
            "total_errores": len(self.errores),
            "total_avisos": len(self.avisos),
            "errores": [asdict(i) for i in self.errores[:MAX_INCIDENCIAS]],
            "avisos": [asdict(i) for i in self.avisos[:MAX_INCIDENCIAS]],
        }


def leer(contenido: bytes) -> Libro:

    _comprobar_archivo(contenido)
    valores = formulas = None
    try:

        valores = openpyxl.load_workbook(
            io.BytesIO(contenido), read_only=True, data_only=True
        )
        formulas = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True)
        hoja_valores, hoja_formulas = valores.worksheets[0], formulas.worksheets[0]
        if not isinstance(hoja_valores, ReadOnlyWorksheet) or not isinstance(
            hoja_formulas, ReadOnlyWorksheet
        ):
            raise DatosInvalidos("La primera hoja del libro no es una hoja de datos")
        return _leer_hoja(hoja_valores, hoja_formulas)
    except DatosInvalidos:
        raise
    except Exception as e:

        raise DatosInvalidos(
            "No se pudo leer el archivo: ¿es un Excel .xlsx válido?"
        ) from e
    finally:
        for libro in (valores, formulas):
            if libro is not None:
                libro.close()


def _comprobar_archivo(contenido: bytes) -> None:
    if not DEFUSEDXML:

        raise RuntimeError(
            "defusedxml no está instalado: no se lee ningún Excel sin él"
        )
    if not contenido.startswith(FIRMA_ZIP):
        raise DatosInvalidos("El archivo no es un Excel .xlsx")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as comprimido:
            partes = comprimido.infolist()
    except zipfile.BadZipFile as e:
        raise DatosInvalidos("El archivo no es un Excel .xlsx") from e
    if (
        len(partes) > MAX_PARTES_ZIP
        or sum(p.file_size for p in partes) > MAX_DESCOMPRIMIDO
    ):
        raise DatosInvalidos(
            "El archivo descomprimido es demasiado grande para un catálogo"
        )


@dataclass
class _Columnas:
    titulo: dict[int, str]
    campo: dict[int, str]
    stock: dict[str, int]
    derivadas: set[int]
    stock_total: int | None


def _leer_hoja(
    hoja_valores: ReadOnlyWorksheet, hoja_formulas: ReadOnlyWorksheet
) -> Libro:
    filas_valores = hoja_valores.iter_rows(values_only=True)
    filas_formulas = hoja_formulas.iter_rows(values_only=True)
    encabezados = next(filas_valores, None)
    next(filas_formulas, None)
    if encabezados is None:
        raise DatosInvalidos(
            "La hoja está vacía: la primera fila debe tener los encabezados"
        )
    columnas, desconocidas = _columnas(encabezados)
    libro = Libro(almacenes=list(columnas.stock))
    for titulo in desconocidas:
        libro.avisos.append(
            Incidencia(1, titulo, "Columna que no es de la plantilla: se ignora")
        )

    fila_de_sku: dict[str, int] = {}
    pares = zip(filas_valores, filas_formulas, strict=True)
    for numero, (valores, formulas) in enumerate(pares, start=2):
        if numero > MAX_FILAS + 1:
            raise DatosInvalidos(
                f"El archivo pasa de {MAX_FILAS} filas: pártelo en varios"
            )
        fila = _leer_fila(numero, valores, formulas, columnas, libro)
        if fila is None:
            continue
        anterior = fila_de_sku.setdefault(fila.datos.sku, numero)
        if anterior != numero:
            libro.errores.append(
                Incidencia(
                    numero,
                    "SKU",
                    f"Repite el SKU {fila.datos.sku} de la fila {anterior}",
                )
            )
            continue
        libro.filas.append(fila)
    return libro


def _columnas(encabezados: tuple) -> tuple[_Columnas, list[str]]:
    por_clave = {_clave(e): e for e in COLUMNAS}
    derivadas = {_clave(e) for e in DERIVADAS}
    columnas = _Columnas(
        titulo={}, campo={}, stock={}, derivadas=set(), stock_total=None
    )
    desconocidas: list[str] = []
    vistas: set[str] = set()
    for i, celda in enumerate(encabezados):
        titulo = " ".join(str(celda).split()) if celda is not None else ""
        if not titulo:
            continue
        clave = _clave(titulo)
        if clave in vistas:
            raise DatosInvalidos(f"La columna «{titulo}» está repetida")
        vistas.add(clave)
        columnas.titulo[i] = titulo
        if clave in por_clave:
            columnas.campo[i] = COLUMNAS[por_clave[clave]]
        elif clave in derivadas:
            columnas.derivadas.add(i)
            if clave == _clave(STOCK_TOTAL):
                columnas.stock_total = i
        elif clave.startswith(_clave(PREFIJO_STOCK) + " "):
            columnas.stock[titulo[len(PREFIJO_STOCK) :].strip()] = i
        else:
            desconocidas.append(titulo)
    presentes = set(columnas.campo.values())
    faltan = [
        e for e, c in COLUMNAS.items() if e not in OPCIONALES and c not in presentes
    ]
    if faltan:
        raise DatosInvalidos(
            f"Faltan columnas de la plantilla: {', '.join(faltan)}. Descarga la plantilla "
            "y copia tus datos en ella."
        )
    return columnas, desconocidas


def _leer_fila(
    numero: int, valores: tuple, formulas: tuple, columnas: _Columnas, libro: Libro
) -> FilaLeida | None:
    def celda(i: int) -> Any:
        return valores[i] if i < len(valores) else None

    if all(_vacia(celda(i)) for i in columnas.campo):
        return None
    errores_antes = len(libro.errores)

    def error(columna: str, mensaje: str) -> None:
        libro.errores.append(Incidencia(numero, columna, mensaje))

    for i, titulo in columnas.titulo.items():
        formula = formulas[i] if i < len(formulas) else None
        if i not in columnas.derivadas and _es_formula(formula) and celda(i) is None:
            error(
                titulo,
                "Tiene una fórmula sin resultado guardado: abre el archivo en Excel y guárdalo",
            )

    datos: dict[str, Any] = {}
    for i, campo in columnas.campo.items():
        titulo = columnas.titulo[i]
        if campo in _DINERO:
            valor = _dinero(celda(i), numero, titulo, libro)
        else:
            valor = _texto(celda(i))
            if campo == "category" and valor is not None:
                valor = _CATEGORIAS.get(_clave(valor), valor)
        if valor is not None:
            datos[campo] = valor

    stock = {
        almacen: _unidades(celda(i), numero, columnas.titulo[i], libro)
        for almacen, i in columnas.stock.items()
    }
    _avisar_si_el_total_no_cuadra(numero, celda, columnas, stock, libro)

    try:
        fila = FilaDeImportacion.model_validate(datos)
    except ValidationError as e:
        libro.errores.extend(_incidencias(numero, e))
        return None
    if len(libro.errores) > errores_antes:
        return None
    return FilaLeida(numero, fila, stock)


def _avisar_si_el_total_no_cuadra(
    numero: int, celda, columnas: _Columnas, stock: dict[str, int], libro: Libro
) -> None:

    if columnas.stock_total is None or not stock:
        return
    total = celda(columnas.stock_total)
    if isinstance(total, bool) or not isinstance(total, int | float):
        return
    suma = sum(stock.values())
    if total != suma:
        libro.avisos.append(
            Incidencia(
                numero,
                STOCK_TOTAL,
                f"Dice {_numero(total)}, pero {' + '.join(stock)} suman {suma}: se carga {suma}",
            )
        )


def _vacia(valor: Any) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _es_formula(valor: Any) -> bool:
    return isinstance(valor, str) and valor.startswith("=")


def _numero(valor: int | float) -> str:
    return (
        str(int(valor))
        if isinstance(valor, float) and valor.is_integer()
        else str(valor)
    )


def _texto(valor: Any) -> str | None:

    if _vacia(valor):
        return None
    texto = (
        _numero(valor)
        if isinstance(valor, int | float) and not isinstance(valor, bool)
        else str(valor)
    )
    return texto.strip() or None


def _dinero(valor: Any, numero: int, columna: str, libro: Libro) -> Decimal | None:
    if _vacia(valor):
        return None
    try:
        if isinstance(valor, bool):
            raise InvalidOperation

        cantidad = Decimal(str(valor).strip())
        redondeada = cantidad.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        libro.errores.append(Incidencia(numero, columna, "No es un número"))
        return None
    if redondeada != cantidad:
        libro.avisos.append(
            Incidencia(
                numero, columna, f"Se redondea {cantidad} a {redondeada} (2 decimales)"
            )
        )
    return redondeada


def _unidades(valor: Any, numero: int, columna: str, libro: Libro) -> int:
    if _vacia(valor):
        return 0
    unidades: int | None = None
    if isinstance(valor, bool):
        pass
    elif isinstance(valor, int):
        unidades = valor
    elif isinstance(valor, float) and valor.is_integer():
        unidades = int(valor)
    elif isinstance(valor, str) and valor.strip().lstrip("-").isdigit():
        unidades = int(valor.strip())
    if unidades is None:
        mensaje = "Debe ser un número entero de unidades"
    elif unidades < 0:
        mensaje = "No puede ser negativo"
    elif unidades > MAX_STOCK_INICIAL:
        mensaje = f"Máximo {MAX_STOCK_INICIAL:,} unidades".replace(",", " ")
    else:
        return unidades
    libro.errores.append(Incidencia(numero, columna, mensaje))
    return 0


_MENSAJES = {
    "missing": "No puede quedar vacío",
    "string_too_short": "No puede quedar vacío",
    "greater_than_equal": "No puede ser negativo",
    "decimal_max_digits": "El número es demasiado grande",
    "enum": "No es una categoría del catálogo",
}


def _incidencias(numero: int, error: ValidationError) -> list[Incidencia]:
    incidencias = []
    for detalle in error.errors(include_url=False):
        campo = detalle["loc"][0] if detalle["loc"] else None
        columna = _ENCABEZADO_DE_CAMPO.get(str(campo)) if campo is not None else None
        if detalle["type"] == "string_too_long":
            mensaje = f"Máximo {detalle.get('ctx', {}).get('max_length')} caracteres"
        else:
            mensaje = _MENSAJES.get(detalle["type"], detalle["msg"])
        incidencias.append(Incidencia(numero, columna, mensaje))
    return incidencias


def revisar(libro: Libro, crear_almacenes: bool = False) -> Revision:

    revision = Revision(errores=list(libro.errores), avisos=list(libro.avisos))
    _resolver_almacenes(libro, revision, crear_almacenes)

    skus = [f.datos.sku for f in libro.filas]
    existentes = {
        p.sku: p
        for p in db.session.execute(
            db.select(DimProducto).where(DimProducto.sku.in_(skus))
        ).scalars()
    }
    activos = [p.id for p in existentes.values() if p.is_active]
    con_historia = productos_service.con_movimientos(activos)
    stock_actual = {
        (s.producto_id, s.almacen_id): s.stock_actual
        for s in db.session.execute(
            db.select(FactStockAlmacen).where(FactStockAlmacen.producto_id.in_(activos))
        ).scalars()
    }

    for fila in libro.filas:
        producto = existentes.get(fila.datos.sku)
        if producto is None:
            _revisar_alta(fila, revision)
        elif not producto.is_active:
            revision.errores.append(
                Incidencia(
                    fila.numero, "SKU", "Es de un producto eliminado: usa otro SKU"
                )
            )
        else:
            _revisar_cambio(
                fila, producto, producto.id in con_historia, stock_actual, revision
            )
    return revision


def _resolver_almacenes(libro: Libro, revision: Revision, crear: bool) -> None:
    activos = {
        _clave(a.nombre): a
        for a in db.session.execute(
            db.select(DimAlmacen).where(DimAlmacen.is_active.is_(True))
        ).scalars()
    }
    for nombre in libro.almacenes:
        columna = f"{PREFIJO_STOCK}{nombre}"
        almacen = activos.get(_clave(nombre))
        if almacen is not None:
            revision.almacenes[nombre] = almacen
        elif crear:
            revision.almacenes_nuevos.append(nombre)
            revision.avisos.append(
                Incidencia(1, columna, f"Se creará el almacén «{nombre}»")
            )
        else:
            revision.errores.append(
                Incidencia(1, columna, f"No hay un almacén activo llamado «{nombre}»")
            )


def _revisar_alta(fila: FilaLeida, revision: Revision) -> None:
    if fila.datos.purchase_price is None:
        revision.errores.append(
            Incidencia(
                fila.numero, "Costo Unitario", "Un producto nuevo necesita su costo"
            )
        )
        return
    valores = fila.datos.model_dump()
    if valores["sale_price"] is None:
        valores["sale_price"] = Decimal("0.00")
    try:
        datos = ProductoCrear.model_validate(valores)
    except ValidationError as e:
        revision.errores.extend(_incidencias(fila.numero, e))
        return
    revision.altas.append(_Alta(fila, datos))


def _revisar_cambio(
    fila: FilaLeida,
    producto: DimProducto,
    tiene_historia: bool,
    stock_actual: dict[tuple[UUID, UUID], int],
    revision: Revision,
) -> None:
    campos: dict[str, Any] = {}
    for campo, valor in fila.datos.model_dump().items():
        if campo in _DINERO and valor is None:
            continue
        if campo == "purchase_price" and tiene_historia:
            continue
        if getattr(producto, COLUMNA_DE_CAMPO[campo]) != valor:
            campos[campo] = valor

    cargar_stock = any(fila.stock.values())
    if tiene_historia:
        cargar_stock = False
        _avisar_lo_que_lleva_el_kardex(fila, producto, stock_actual, revision)

    if campos or cargar_stock:
        revision.cambios.append(_Cambio(fila, producto, campos, cargar_stock))
    else:
        revision.sin_cambios += 1


def _avisar_lo_que_lleva_el_kardex(
    fila: FilaLeida,
    producto: DimProducto,
    stock_actual: dict[tuple[UUID, UUID], int],
    revision: Revision,
) -> None:
    costo = fila.datos.purchase_price
    if costo is not None and costo != producto.precio_compra_actual:
        revision.avisos.append(
            Incidencia(
                fila.numero,
                "Costo Unitario",
                "Tiene movimientos: el costo lo lleva el Kardex "
                f"(S/ {producto.precio_compra_actual}) y no se cambia",
            )
        )
    for nombre, unidades in fila.stock.items():
        almacen = revision.almacenes.get(nombre)
        actual = stock_actual.get((producto.id, almacen.id), 0) if almacen else 0
        if unidades != actual:
            revision.avisos.append(
                Incidencia(
                    fila.numero,
                    f"{PREFIJO_STOCK}{nombre}",
                    f"Tiene movimientos: el stock lo lleva el Kardex ({actual}) y no se cambia",
                )
            )


def plantilla() -> bytes:
    almacenes = db.session.execute(
        db.select(DimAlmacen.nombre)
        .where(DimAlmacen.is_active.is_(True))
        .order_by(DimAlmacen.nombre)
    ).scalars()
    encabezados = list(COLUMNAS)
    corte = encabezados.index("Costo Unitario")
    encabezados[corte:corte] = [f"{PREFIJO_STOCK}{nombre}" for nombre in almacenes]

    libro = openpyxl.Workbook()
    libro.remove(libro.worksheets[0])
    hoja = libro.create_sheet("Productos")
    hoja.append(encabezados)
    hoja.freeze_panes = "A2"
    for i, encabezado in enumerate(encabezados, start=1):
        letra = get_column_letter(i)
        hoja.cell(row=1, column=i).font = Font(bold=True)
        hoja.column_dimensions[letra].width = max(14, len(encabezado) + 4)
        if encabezado in ("SKU", "Código Barras"):
            hoja.column_dimensions[letra].number_format = "@"

    listas = libro.create_sheet("Listas")
    listas.sheet_state = "hidden"
    for fila, categoria in enumerate(CategoriaProducto, start=1):
        listas.cell(row=fila, column=1, value=categoria.value)
    ultima = f"{MAX_FILAS + 1}"
    validaciones = [
        (
            DataValidation(
                type="list", formula1=f"=Listas!$A$1:$A${len(CategoriaProducto)}"
            ),
            ["Categoria"],
            "Elige una categoría de la lista",
        ),
        (
            DataValidation(
                type="whole",
                operator="between",
                formula1="0",
                formula2=f"{MAX_STOCK_INICIAL}",
            ),
            [e for e in encabezados if e.startswith(PREFIJO_STOCK)],
            "Unidades enteras, 0 o más",
        ),
        (
            DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0"),
            ["Costo Unitario", "Precio Venta"],
            "Un importe en soles, 0 o más",
        ),
    ]
    for validacion, columnas, mensaje in validaciones:
        validacion.allow_blank = True
        validacion.showErrorMessage = True
        validacion.error = mensaje
        for encabezado in columnas:
            letra = get_column_letter(encabezados.index(encabezado) + 1)
            validacion.add(f"{letra}2:{letra}{ultima}")
        hoja.add_data_validation(validacion)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def aplicar(revision: Revision, usuario_id: UUID) -> None:
    if revision.errores:
        raise DatosInvalidos(
            f"El archivo tiene {len(revision.errores)} errores: no se aplicó nada"
        )
    with productos_service.transaccion():
        for nombre in revision.almacenes_nuevos:
            almacen = DimAlmacen()
            almacen.nombre = nombre
            db.session.add(almacen)
            revision.almacenes[nombre] = almacen
        db.session.flush()
        for alta in revision.altas:
            producto = DimProducto()
            for campo, valor in alta.datos.model_dump(
                exclude={"initial_stock", "warehouse_id"}
            ).items():
                setattr(producto, COLUMNA_DE_CAMPO[campo], valor)
            db.session.add(producto)
            productos_service.escribir_sin_duplicado(producto.sku)
            _cargar_stock(producto, alta.fila, revision, usuario_id)
        for cambio in revision.cambios:
            for campo, valor in cambio.campos.items():
                setattr(cambio.producto, COLUMNA_DE_CAMPO[campo], valor)
            productos_service.escribir_sin_duplicado(cambio.producto.sku)
            if cambio.cargar_stock:
                _cargar_stock(cambio.producto, cambio.fila, revision, usuario_id)


def _cargar_stock(
    producto: DimProducto, fila: FilaLeida, revision: Revision, usuario_id: UUID
) -> None:
    for nombre, unidades in fila.stock.items():
        if unidades == 0:
            continue
        carga = CargaInicial(
            initial_stock=unidades, warehouse_id=revision.almacenes[nombre].id
        )
        productos_service.cargar_stock_inicial(
            producto, carga, usuario_id, origen="EXCEL"
        )
